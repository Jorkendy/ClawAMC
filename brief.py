"""Buoc 5 — Brief Design generator. build_brief_content (enrich LLM) + render_brief_pptx (python-pptx).

AI nhap -> M&D hoan thien. Chu dong fill gap khi input ngheo + GAN NHAN nguon (requester|ai)
de M&D biet cai nao la gia dinh. Khong bia chi tiet logo/KV brand cu the.
"""
import base64
import concurrent.futures
import io

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from llm_client import ask_llm_json, generate_image
from proposal import game_insight

BRIEF_PROMPT = """Bạn là chuyên gia brief design merchandise game. Soạn BRIEF DESIGN cho bộ quà dưới đây \
để đội M&D/đối tác design. Đây là BẢN NHÁP định hướng — bạn CHỦ ĐỘNG đề xuất khi thiếu thông tin, \
nhưng PHẢI gắn nhãn nguồn từng phần: "requester" (lấy từ dữ liệu requester cung cấp) hoặc "ai" (bạn suy luận/đề xuất).

QUY TẮC:
- Định hướng visual suy từ INSIGHT GAME + loại item. TUYỆT ĐỐI KHÔNG bịa chi tiết logo/KV/brand cụ thể \
(màu mã hex thật, vị trí logo chính xác) — chỉ gợi ý hướng, đánh "ai".
- Trường nào requester đã ghi rõ (chất liệu, kích thước, yêu cầu đặc biệt) → giữ nguyên, đánh "requester".
- "print_spec": gợi ý vùng in / số màu / định dạng file cần (vd "file vector AI", "file mockup") theo loại item.
- "references": gom link requester đưa (trong yêu cầu/design_link) + loại tài liệu tham khảo nên có (đánh "ai").
- "image_prompt": MÔ TẢ TIẾNG ANH ngắn (1 câu) sản phẩm vật lý cho công cụ sinh ảnh AI. PHẢI gọi đúng TÊN \
sản phẩm bằng tiếng Anh (vd "acrylic keychain", "hardcover notebook with elastic band", "insulated stainless \
steel water bottle", "enamel pin badge", "canvas tote bag") + hình dáng/chất liệu/màu chủ đạo. KHÔNG mô tả chữ/logo \
cụ thể (bề mặt sản phẩm sẽ để trống, ảnh không render chữ).
- Các trường khác Tiếng Việt, ngắn gọn, đúng trọng tâm cho designer.

THÔNG TIN BỘ QUÀ:
- Game: {game}
- Mục đích: {muc_dich}
- Chủ đề: {chu_de}
- Đối tượng: {target}
- Yêu cầu đặc biệt chung: {yeu_cau}
- Trạng thái asset: {asset_status}

INSIGHT GAME (định hướng visual khách quan):
{insight}

CÁC ITEM ĐÃ CHỐT:
{items_block}

JSON schema (CHỈ trả JSON, không text thừa):
{{"collection_name": str, "overview": str, "asset_status": str, "items": [{{"ten": str, "loai": str, \
"idea": str, "design_direction": str, "chat_lieu": str, "kich_thuoc": str, "yeu_cau_dac_biet": [str], \
"print_spec": str, "image_prompt": str, "references": [str], "sources": {{"idea": "requester"|"ai", "design_direction": "requester"|"ai", \
"chat_lieu": "requester"|"ai", "kich_thuoc": "requester"|"ai", "print_spec": "requester"|"ai"}}}}]}}"""


def _items_block(items: list) -> str:
    lines = []
    for it in items:
        key = " ⭐" if it.get("item_key") else ""
        lines.append(
            f"- {it.get('ten','')}{key} | loại: {it.get('loai','')} | nguồn: {it.get('nguon','')} "
            f"| chất liệu: {it.get('chat_lieu') or '(chưa có)'} | kích thước: {it.get('kich_thuoc') or '(chưa có)'} "
            f"| SL: {it.get('so_luong') or '?'} | yêu cầu: {it.get('yeu_cau_dac_biet') or '(không)'} "
            f"| design link: {it.get('design_link') or '(không)'}")
    return "\n".join(lines)


def _asset_status_text(asset_status: dict) -> str:
    parts = []
    parts.append("Logo: " + ("đã có" if asset_status.get("logo") else "CHƯA có"))
    parts.append("KV: " + ("đã có" if asset_status.get("kv") else "CHƯA có"))
    parts.append("Source: " + ("đã có" if asset_status.get("source") else "CHƯA có"))
    return " · ".join(parts)


def build_brief_content(fields: dict, items: list, asset_status: dict, insight: str) -> dict:
    """Goi 1 LLM call enrich noi dung brief. Tra dict theo schema (xem docstring module/spec)."""
    prompt = BRIEF_PROMPT.format(
        game=fields.get("Game") or "(chưa rõ)",
        muc_dich=fields.get("Mục đích") or "(chưa rõ)",
        chu_de=fields.get("Chủ đề") or "(chưa rõ)",
        target=fields.get("Target audience") or "(chưa rõ)",
        yeu_cau=fields.get("Yêu cầu đặc biệt") or "(không)",
        asset_status=_asset_status_text(asset_status),
        insight=insight or "(không có insight)",
        items_block=_items_block(items),
    )
    data = ask_llm_json(prompt, max_tokens=3000)
    # chuan hoa toi thieu (chong thieu key lam vo render)
    data.setdefault("collection_name", fields.get("Chủ đề") or "Bộ quà merch")
    data.setdefault("overview", "")
    data.setdefault("asset_status", _asset_status_text(asset_status))
    out_items = []
    for it in data.get("items", []):
        it.setdefault("ten", "")
        it.setdefault("loai", "")
        for k in ("idea", "design_direction", "chat_lieu", "kich_thuoc", "print_spec", "image_prompt"):
            it.setdefault(k, "")
        if not isinstance(it.get("yeu_cau_dac_biet"), list):
            it["yeu_cau_dac_biet"] = []
        if not isinstance(it.get("references"), list):
            it["references"] = []
        if not isinstance(it.get("sources"), dict):
            it["sources"] = {}
        out_items.append(it)
    data["items"] = out_items
    return data


# ---------- gather images (tai dung cache proposal) ----------
_BRIEF_IMG_CACHE: dict = {}
_BRIEF_IMG_TIMEOUT = 120  # giay/anh (cap CUNG nho max_retries=0) — latency anh dao dong ~35-90s, het gio bo qua (chan treo pipeline)
_BRIEF_IMG_WORKERS = 3    # gen SONG SONG (giam wall-clock; nhe tay voi endpoint)


def _brief_img_prompt(it: dict, game: str) -> str:
    # Uu tien image_prompt (TIENG ANH, do LLM enrich sinh) — Imagen ve san pham ro hon han ten tieng Viet.
    # Fallback khi thieu: ghep loai+ten+direction (kem hon nhung con hon khong).
    desc = (it.get("image_prompt") or "").strip()
    if not desc:
        direction = (it.get("design_direction") or it.get("idea") or "").strip()
        desc = f"{it.get('loai', '')} {it.get('ten', '')} — {direction}"
    return (f"Photorealistic studio product photograph of a single physical product: {desc}. "
            f"Merchandise concept inspired by the video game '{game}'. "
            "Centered on a plain neutral background, soft even lighting, clean concept mockup. "
            "IMPORTANT: the product surface must be BLANK — do NOT render any text, letters, words, "
            "numbers, captions, labels or typography anywhere in the image.")


def _b64_to_bytes(b64: str | None) -> bytes | None:
    """Decode base64 -> PNG bytes; None/loi decode -> None."""
    if not b64:
        return None
    try:
        return base64.b64decode(b64)
    except Exception:
        return None


def gather_brief_images(brief_items: list, game: str) -> dict:
    """Map {ten item -> PNG bytes}: AI generate concept mockup theo GAME cho MOI item
    (catalogue + creative), dua tren design_direction da enrich — KHONG dung anh catalogue goc.
    Gen SONG SONG + timeout/anh: tranh treo ca pipeline (truoc tuan tu, 1 call ket -> ket het).
    Cache theo prompt chuan hoa -> re-trigger khong gen lai. Timeout/loi 1 anh -> bo qua item."""
    out, todo = {}, []
    for it in brief_items:
        name = (it.get("ten") or "").strip()
        if not name:
            continue
        prompt = _brief_img_prompt(it, game)
        key = " ".join(prompt.lower().split())
        cached = _BRIEF_IMG_CACHE.get(key)
        if cached:
            img = _b64_to_bytes(cached)
            if img:
                out[name] = img
        else:
            todo.append((name, key, prompt))
    if not todo:
        return out
    with concurrent.futures.ThreadPoolExecutor(max_workers=_BRIEF_IMG_WORKERS) as ex:
        futs = {ex.submit(generate_image, p, _BRIEF_IMG_TIMEOUT): (n, k) for (n, k, p) in todo}
        for f in concurrent.futures.as_completed(futs):
            name, key = futs[f]
            try:
                b64 = f.result()
            except Exception:
                b64 = None
            if not b64:
                continue
            _BRIEF_IMG_CACHE[key] = b64  # chi cache khi gen thanh cong
            img = _b64_to_bytes(b64)
            if img:
                out[name] = img
    return out


# ---------- render pptx ----------
_C_BG = RGBColor(0x1A, 0x1A, 0x2E)
_C_ACCENT = RGBColor(0xFF, 0x3D, 0x57)
_C_REQ = RGBColor(0x1A, 0x7F, 0x37)     # xanh = requester
_C_AI = RGBColor(0xD9, 0x77, 0x06)      # cam = AI goi y
_C_TEXT = RGBColor(0x1A, 0x1A, 0x2E)
_C_MUTE = RGBColor(0x6B, 0x72, 0x80)

# python-pptx magic number -> dat ten cho de doc (gia tri KHONG doi)
_BLANK_LAYOUT = 6   # slide_layouts[6] = layout "Blank"
_SHAPE_RECT = 1     # add_shape(1) = MSO_AUTO_SHAPE_TYPE.RECTANGLE


def _tag(source: str) -> tuple:
    return ("(Từ requester)", _C_REQ) if source == "requester" else ("(AI gợi ý)", _C_AI)


def _section(tf, heading: str, body: str, source: str | None, first: bool = False):
    """Them 1 muc vao text frame: heading dam + tag nguon mau + body."""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    r = p.add_run(); r.text = heading + "  "
    r.font.bold = True; r.font.size = Pt(12); r.font.color.rgb = _C_TEXT
    if source:
        label, color = _tag(source)
        rt = p.add_run(); rt.text = label
        rt.font.size = Pt(9); rt.font.color.rgb = color; rt.font.italic = True
    pb = tf.add_paragraph()
    rb = pb.add_run(); rb.text = body or "—"
    rb.font.size = Pt(11); rb.font.color.rgb = _C_TEXT
    pb.space_after = Pt(6)


def _cover(prs, brief_data: dict, project: dict):
    s = prs.slides.add_slide(prs.slide_layouts[_BLANK_LAYOUT])
    band = s.shapes.add_shape(_SHAPE_RECT, 0, 0, prs.slide_width, Inches(2.2))
    band.fill.solid(); band.fill.fore_color.rgb = _C_BG; band.line.fill.background()
    tf = band.text_frame; tf.word_wrap = True
    tf.margin_left = Inches(0.6); tf.margin_top = Inches(0.4)
    p = tf.paragraphs[0]; r = p.add_run(); r.text = brief_data.get("collection_name", "Brief Design")
    r.font.size = Pt(30); r.font.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p2 = tf.add_paragraph(); r2 = p2.add_run()
    r2.text = f"BRIEF DESIGN · {project.get('game','')} · {project.get('code','')}"
    r2.font.size = Pt(13); r2.font.color.rgb = RGBColor(0xFF, 0xC1, 0xC9)

    box = s.shapes.add_textbox(Inches(0.6), Inches(2.5), Inches(12), Inches(4.5))
    tf2 = box.text_frame; tf2.word_wrap = True
    info = [
        ("Tổng quan", brief_data.get("overview", "")),
        ("Thông tin", f"Số lượng: {project.get('so_luong','?')} · Deadline: {project.get('deadline','?')}"),
        ("Trạng thái asset", brief_data.get("asset_status", "")),
    ]
    first = True
    for h, b in info:
        _section(tf2, h, b, None, first=first); first = False
    pl = tf2.add_paragraph(); pl.space_before = Pt(10)
    rl = pl.add_run(); rl.text = "Chú giải:  "
    rl.font.size = Pt(10); rl.font.bold = True
    r_req = pl.add_run(); r_req.text = "(Từ requester) "; r_req.font.size = Pt(10); r_req.font.color.rgb = _C_REQ
    r_ai = pl.add_run(); r_ai.text = " (AI gợi ý)"; r_ai.font.size = Pt(10); r_ai.font.color.rgb = _C_AI
    if project.get("logo_png"):
        try:
            s.shapes.add_picture(io.BytesIO(project["logo_png"]), Inches(10.8), Inches(0.4),
                                 height=Inches(1.3))
        except Exception:
            pass


def _item_slide(prs, it: dict, img: bytes | None):
    s = prs.slides.add_slide(prs.slide_layouts[_BLANK_LAYOUT])
    bar = s.shapes.add_shape(_SHAPE_RECT, 0, 0, prs.slide_width, Inches(0.9))
    bar.fill.solid(); bar.fill.fore_color.rgb = _C_ACCENT; bar.line.fill.background()
    tf = bar.text_frame; tf.margin_left = Inches(0.5)
    p = tf.paragraphs[0]; r = p.add_run()
    r.text = f"{it.get('ten','')}   ·   {it.get('loai','')}"
    r.font.size = Pt(20); r.font.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    has_img = bool(img)
    text_w = Inches(7.4) if has_img else Inches(12.2)
    box = s.shapes.add_textbox(Inches(0.5), Inches(1.2), text_w, Inches(5.9))
    bt = box.text_frame; bt.word_wrap = True
    src = it.get("sources", {})
    _section(bt, "IDEA", it.get("idea", ""), src.get("idea"), first=True)
    _section(bt, "ĐỊNH HƯỚNG DESIGN", it.get("design_direction", ""), src.get("design_direction"))
    _section(bt, "Chất liệu", it.get("chat_lieu", ""), src.get("chat_lieu"))
    _section(bt, "Kích thước", it.get("kich_thuoc", ""), src.get("kich_thuoc"))
    special_reqs = it.get("yeu_cau_dac_biet") or []
    _section(bt, "Yêu cầu đặc biệt", ("; ".join(special_reqs) if special_reqs else "Không"), None)
    _section(bt, "Print spec / file cần", it.get("print_spec", ""), src.get("print_spec"))
    refs = it.get("references") or []
    _section(bt, "References", ("\n".join(refs) if refs else "—"), None)

    if has_img:
        try:
            s.shapes.add_picture(io.BytesIO(img), Inches(8.1), Inches(1.3), width=Inches(4.7))
            cap = s.shapes.add_textbox(Inches(8.1), Inches(6.1), Inches(4.7), Inches(0.5))
            rc = cap.text_frame.paragraphs[0].add_run()
            rc.text = "Hình minh hoạ ý tưởng — AI tạo, chưa phải mẫu cuối"
            rc.font.size = Pt(8); rc.font.italic = True; rc.font.color.rgb = _C_MUTE
        except Exception:
            pass


def render_brief_pptx(brief_data: dict, images: dict, project: dict) -> bytes:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    _cover(prs, brief_data, project)
    for it in brief_data.get("items", []):
        _item_slide(prs, it, (images or {}).get(it.get("ten", "")))
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
