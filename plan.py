"""Buoc 4 — Plan san xuat (deterministic, KHONG dung LLM).

Input: items da chot (sau Duyet) + catalogue lead-time + don gia da chot.
Output: file Excel (.xlsx) gom:
  - Sheet "Ke hoach tong": timeline overall theo critical path (san xuat song song -> max),
    cac giai doan kem ngay bat dau/ket thuc + so deadline kha thi/tre + phan bo ngan sach.
  - Sheet "Chi tiet items": tung item (so luong, don gia, thanh tien, len mau/san xuat, ghi chu).

Moi con so suy tu catalogue + don gia da chot — khong nho LLM (LLM cong sai + tinh ngay kem).
"""
import io
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from config import (OVERHEAD_DELIVERY_WORKDAYS, OVERHEAD_HEAD_WORKDAYS,
                    OVERHEAD_REVIEW_SAMPLE_WORKDAYS, WORKDAYS_TO_CALENDAR)
from proposal import catalogue_data, deadline_days_needed, item_leadtime

ADVANCE_RATIO = 0.5  # [GIA DINH] tam ung 50% khi duyet mau, 50% khi nghiem thu — can xac nhan dieu khoan


def _wd_to_cal(workdays: int) -> int:
    return round(workdays * WORKDAYS_TO_CALENDAR)


def _fmt_d(d: date) -> str:
    return d.strftime("%d/%m/%Y")


def _parse_deadline(s) -> date | None:
    if not s:
        return None
    try:
        return date.fromisoformat(str(s)[:10])
    except Exception:
        return None


def build_plan(fields: dict, items: list, start: date) -> dict:
    """Tinh toan plan (thuan, de test). Tra dict moi gia tri da tinh san de render."""
    _, by_name = catalogue_data()
    code = fields.get("Mã project") or ""
    name = fields.get("Tên project") or ""
    qty = fields.get("Số lượng (bộ/suất)") or 0
    budget = int(fields.get("Budget (VND)") or 0)
    deadline = _parse_deadline(fields.get("Deadline cần hàng"))

    # --- lead-time tung item + thanh tien ---
    rows, leads = [], []
    for it in items:
        lm, sx = item_leadtime(it, by_name)
        leads.append((lm, sx))
        is_cat = it.get("nguon") == "catalogue"
        dg = it.get("don_gia") if is_cat else None
        so_luong = int(it.get("so_luong") or 0)
        thanh_tien = (dg or 0) * so_luong
        rows.append({
            "ten": it.get("ten") or "", "loai": it.get("loai") or "",
            "nguon": "Catalogue" if is_cat else "Creative",
            "so_luong": so_luong, "don_gia": dg, "thanh_tien": thanh_tien,
            "len_mau": lm, "san_xuat": sx,
            "ghi_chu": "" if is_cat else "Giá & timeline DỰ KIẾN — chốt sau khi có design + báo giá vendor",
        })
    max_lm = max((lm for lm, _ in leads), default=0)
    max_sx = max((sx for _, sx in leads), default=0)

    # --- timeline tong theo critical path (overhead 27 = head 18 + duyet mau 7 + giao 2) ---
    phases_wd = [
        ("Chuẩn bị (brief, kế hoạch, chốt NCC)", OVERHEAD_HEAD_WORKDAYS),
        ("Lên mẫu", max_lm),
        ("Duyệt mẫu", OVERHEAD_REVIEW_SAMPLE_WORKDAYS),
        ("Sản xuất", max_sx),
        ("Giao hàng & nghiệm thu", OVERHEAD_DELIVERY_WORKDAYS),
    ]
    total_cal = deadline_days_needed(items, by_name)  # nguon chuan (khop check kha thi o noi khac)
    cals = [_wd_to_cal(wd) for _, wd in phases_wd]
    cals[-1] += total_cal - sum(cals)  # phase cuoi hap thu sai so lam tron -> cong don = total_cal

    phases, cursor = [], start
    for (label, wd), cal in zip(phases_wd, cals):
        ph_start = cursor
        ph_end = cursor + timedelta(days=max(cal, 0))
        phases.append({"label": label, "workdays": wd, "cal_days": cal,
                       "start": ph_start, "end": ph_end})
        cursor = ph_end
    delivery = start + timedelta(days=total_cal)

    # cua so len mau / san xuat (de ve checklist tung item)
    len_mau_phase = next(p for p in phases if p["label"] == "Lên mẫu")
    san_xuat_phase = next(p for p in phases if p["label"] == "Sản xuất")
    for r in rows:
        r["len_mau_start"] = len_mau_phase["start"]
        r["len_mau_end"] = len_mau_phase["start"] + timedelta(days=_wd_to_cal(r["len_mau"]))
        r["san_xuat_start"] = san_xuat_phase["start"]
        r["san_xuat_end"] = san_xuat_phase["start"] + timedelta(days=_wd_to_cal(r["san_xuat"]))

    # --- kha thi deadline ---
    if deadline is None:
        feasible = "(không có deadline để đối chiếu)"
    elif delivery <= deadline:
        feasible = f"✅ Khả thi — dư {(deadline - delivery).days} ngày"
    else:
        feasible = f"🔴 TRỄ {(delivery - deadline).days} ngày so với deadline"

    # --- phan bo ngan sach ---
    tong_co_gia = sum(r["thanh_tien"] for r in rows if r["don_gia"])
    has_creative_no_price = any(r["nguon"] == "Creative" for r in rows)
    advance = round(tong_co_gia * ADVANCE_RATIO)

    return {
        "code": code, "name": name, "qty": qty, "budget": budget,
        "deadline": _fmt_d(deadline) if deadline else "(chưa có)",
        "start": _fmt_d(start), "delivery": _fmt_d(delivery),
        "total_cal": total_cal, "feasible": feasible,
        "phases": phases, "rows": rows,
        "tong_co_gia": tong_co_gia, "con_lai": budget - tong_co_gia,
        "has_creative_no_price": has_creative_no_price,
        "advance": advance, "final_pay": tong_co_gia - advance,
    }


# ---------- render Excel ----------
_HEAD_FILL = PatternFill("solid", fgColor="305496")
_SUB_FILL = PatternFill("solid", fgColor="D9E1F2")
_HEAD_FONT = Font(bold=True, color="FFFFFF")
_BOLD = Font(bold=True)
_THIN = Side(style="thin", color="BFBFBF")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_MONEY = '#,##0" đ"'


def _money(cell):
    cell.number_format = _MONEY
    cell.alignment = Alignment(horizontal="right")


def _header_row(ws, row, headers, start_col=1):
    for i, h in enumerate(headers):
        c = ws.cell(row=row, column=start_col + i, value=h)
        c.fill, c.font, c.border = _HEAD_FILL, _HEAD_FONT, _BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def render_plan_xlsx(plan: dict) -> bytes:
    wb = Workbook()

    # ===== Sheet 1: Ke hoach tong =====
    ws = wb.active
    ws.title = "Kế hoạch tổng"
    ws.column_dimensions["A"].width = 34
    for col in "BCDE":
        ws.column_dimensions[col].width = 18

    t = ws.cell(row=1, column=1, value=f"KẾ HOẠCH SẢN XUẤT — {plan['code']}")
    t.font = Font(bold=True, size=14)
    ws.merge_cells("A1:E1")

    info = [
        ("Tên project", plan["name"]),
        ("Mã project", plan["code"]),
        ("Số lượng (bộ/suất)", plan["qty"]),
        ("Deadline cần hàng", plan["deadline"]),
        ("Ngày bắt đầu (duyệt)", plan["start"]),
        ("Ngày giao dự kiến", plan["delivery"]),
        ("Tổng thời gian dự kiến", f"{plan['total_cal']} ngày lịch"),
        ("Khả thi deadline", plan["feasible"]),
    ]
    r = 3
    for k, v in info:
        ws.cell(row=r, column=1, value=k).font = _BOLD
        ws.cell(row=r, column=2, value=v)
        r += 1

    # Timeline tong
    r += 1
    s = ws.cell(row=r, column=1, value="TIMELINE TỔNG (sản xuất song song theo item → lấy max)")
    s.font = _BOLD
    s.fill = _SUB_FILL
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    r += 1
    _header_row(ws, r, ["Giai đoạn", "Số ngày (LV)", "Số ngày (lịch)", "Bắt đầu", "Kết thúc"])
    r += 1
    for ph in plan["phases"]:
        ws.cell(row=r, column=1, value=ph["label"]).border = _BORDER
        ws.cell(row=r, column=2, value=ph["workdays"]).border = _BORDER
        ws.cell(row=r, column=3, value=ph["cal_days"]).border = _BORDER
        ws.cell(row=r, column=4, value=_fmt_d(ph["start"])).border = _BORDER
        ws.cell(row=r, column=5, value=_fmt_d(ph["end"])).border = _BORDER
        r += 1

    # Phan bo ngan sach
    r += 1
    s = ws.cell(row=r, column=1, value="PHÂN BỔ NGÂN SÁCH")
    s.font = _BOLD
    s.fill = _SUB_FILL
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    r += 1
    budget_rows = [
        ("Tổng chi phí dự kiến (hàng có giá)", plan["tong_co_gia"]),
        ("Budget", plan["budget"]),
        ("Còn lại", plan["con_lai"]),
    ]
    for k, v in budget_rows:
        ws.cell(row=r, column=1, value=k).font = _BOLD
        _money(ws.cell(row=r, column=2, value=v))
        r += 1
    if plan["has_creative_no_price"]:
        ws.cell(row=r, column=1, value="Lưu ý").font = _BOLD
        ws.cell(row=r, column=2,
                value="Còn item creative chưa có giá — chờ báo giá vendor sau khi có design")
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
        r += 1

    # Dot chi (gia dinh)
    r += 1
    s = ws.cell(row=r, column=1, value="ĐỀ XUẤT ĐỢT CHI  [GIẢ ĐỊNH — cần xác nhận điều khoản thanh toán]")
    s.font = Font(bold=True, italic=True)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    r += 1
    for k, v in [("Tạm ứng 50% (khi duyệt mẫu)", plan["advance"]),
                 ("Thanh toán 50% (khi nghiệm thu & giao)", plan["final_pay"])]:
        ws.cell(row=r, column=1, value=k)
        _money(ws.cell(row=r, column=2, value=v))
        r += 1

    # ===== Sheet 2: Chi tiet items =====
    ws2 = wb.create_sheet("Chi tiết items")
    headers = ["STT", "Tên item", "Loại", "Nguồn", "Số lượng", "Đơn giá", "Thành tiền",
               "Lên mẫu (LV)", "Sản xuất (LV)", "Lên mẫu (dự kiến)", "Sản xuất (dự kiến)", "Ghi chú"]
    widths = [5, 26, 16, 11, 9, 14, 16, 11, 12, 22, 22, 40]
    for i, w in enumerate(widths):
        ws2.column_dimensions[get_column_letter(i + 1)].width = w
    _header_row(ws2, 1, headers)
    rr = 2
    for idx, it in enumerate(plan["rows"], 1):
        ws2.cell(row=rr, column=1, value=idx).border = _BORDER
        ws2.cell(row=rr, column=2, value=it["ten"]).border = _BORDER
        ws2.cell(row=rr, column=3, value=it["loai"]).border = _BORDER
        ws2.cell(row=rr, column=4, value=it["nguon"]).border = _BORDER
        ws2.cell(row=rr, column=5, value=it["so_luong"]).border = _BORDER
        c6 = ws2.cell(row=rr, column=6, value=it["don_gia"] if it["don_gia"] else "—")
        c6.border = _BORDER
        if it["don_gia"]:
            _money(c6)
        c7 = ws2.cell(row=rr, column=7, value=it["thanh_tien"] if it["don_gia"] else "—")
        c7.border = _BORDER
        if it["don_gia"]:
            _money(c7)
        ws2.cell(row=rr, column=8, value=it["len_mau"]).border = _BORDER
        ws2.cell(row=rr, column=9, value=it["san_xuat"]).border = _BORDER
        ws2.cell(row=rr, column=10,
                 value=f"{_fmt_d(it['len_mau_start'])} → {_fmt_d(it['len_mau_end'])}").border = _BORDER
        ws2.cell(row=rr, column=11,
                 value=f"{_fmt_d(it['san_xuat_start'])} → {_fmt_d(it['san_xuat_end'])}").border = _BORDER
        gc = ws2.cell(row=rr, column=12, value=it["ghi_chu"])
        gc.border = _BORDER
        gc.alignment = Alignment(wrap_text=True)
        rr += 1
    # tong
    ws2.cell(row=rr, column=2, value="TỔNG").font = _BOLD
    _money(ws2.cell(row=rr, column=7, value=plan["tong_co_gia"]))
    ws2.cell(row=rr, column=7).font = _BOLD

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
